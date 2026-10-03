from abc import ABC, abstractmethod
from io import BytesIO

class TextExtractor(ABC):
    @abstractmethod
    def extract(self, filename: str, data: bytes) -> str:
        raise NotImplementedError

class PlainTextExtractor(TextExtractor):
    def extract(self, filename: str, data: bytes) -> str:
        return data.decode("utf-8", errors="replace")

class PdfExtractor(TextExtractor):
    def extract(self, filename: str, data: bytes) -> str:
        from pypdf import PdfReader
        reader=PdfReader(BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

class DocxExtractor(TextExtractor):
    def extract(self, filename: str, data: bytes) -> str:
        from docx import Document
        doc=Document(BytesIO(data))
        parts=[p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)

class ExtractorRegistry:
    def __init__(self):
        self._extractors={
            "txt":PlainTextExtractor(),
            "md":PlainTextExtractor(),
            "csv":PlainTextExtractor(),
            "pdf":PdfExtractor(),
            "docx":DocxExtractor(),
        }
    def get(self, filename: str) -> TextExtractor:
        suffix=filename.rsplit(".",1)[-1].lower() if "." in filename else "txt"
        if suffix not in self._extractors:
            raise ValueError(f"Unsupported knowledge file type: {suffix}")
        return self._extractors[suffix]
